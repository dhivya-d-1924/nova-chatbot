"""Document ingestion and retrieval.

Extraction: PDF (text layer, OCR fallback for scanned pages), DOCX, TXT/MD/CSV, and images via OCR.
Hardening: content is checked by magic bytes, DOCX zip contents are size-checked (zip bombs), PDF page counts
are capped. Retrieval: BM25 keyword ranking always; when EMBEDDING_MODEL is set, vector similarity is fused
with BM25 using reciprocal rank fusion. Every query is filtered by the owner's user id.
"""
import array
import hashlib
import importlib.util
import io
import logging
import math
import re
import shutil
import uuid
import zipfile
from collections import Counter

from sqlalchemy import delete, func, insert, select, update

from . import db, llm
from .config import settings
from .errors import ApiError

log = logging.getLogger("nova.rag")

TEXT_EXT = {"txt", "md", "csv"}
IMAGE_EXT = {"png", "jpg", "jpeg", "webp", "gif", "tif", "tiff"}
BLOCKED_EXT = {"exe", "dll", "bat", "cmd", "sh", "ps1", "js", "msi", "com", "scr", "jar", "py", "php", "html", "htm", "svg"}
STOPWORDS = set("a an and are as at be by for from has have in is it its of on or that the this to was were will with "
                "what which who how when where why do does did can could should would about into than then there their "
                "you your i me my we our they them he she his her not no".split())


def ocr_available():
    return settings.OCR_ENABLED and bool(importlib.util.find_spec("pytesseract")) and bool(shutil.which("tesseract"))


def allowed_extensions():
    exts = {"pdf", "docx", *TEXT_EXT}
    if ocr_available():
        exts |= IMAGE_EXT
    return sorted(exts)


# ------------------------------------------------------------------ text processing
def tokenize(t):
    return [w for w in re.findall(r"\w+", t.lower()) if len(w) > 1 and w not in STOPWORDS]


def bm25_scores(query, bodies):
    docs = [tokenize(b) for b in bodies]
    n = len(docs)
    if not n:
        return []
    avg = (sum(len(d) for d in docs) / n) or 1
    df = Counter()
    for d in docs:
        df.update(set(d))
    q = set(tokenize(query))
    k1, b = 1.5, 0.75
    scores = []
    for d in docs:
        tf = Counter(d)
        scores.append(sum(math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5)) * tf[t] * (k1 + 1)
                          / (tf[t] + k1 * (1 - b + b * len(d) / avg)) for t in q if t in tf))
    return scores


def split_text(text, size=1000, overlap=150):
    """Split into ~size-char chunks on sentence/word boundaries. Returns [(start_offset, chunk)]."""
    text = re.sub(r"\s+", " ", text).strip()
    out, i = [], 0
    while i < len(text):
        end = min(len(text), i + size)
        if end < len(text):
            cut = max(text.rfind(". ", i, end) + 1, text.rfind(" ", i, end))
            if cut > i + size // 2:
                end = cut
        piece = text[i:end].strip()
        if piece:
            out.append((i, piece))
        if end >= len(text):
            break
        i = max(end - overlap, i + 1)
    return out


# ------------------------------------------------------------------ extraction
def detect_kind(filename, data):
    """Decide the parser from the file's bytes, not just its name."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in BLOCKED_EXT or data[:2] == b"MZ" or data[:4] == b"\x7fELF":
        raise ApiError(400, "Executable or script files can't be uploaded.", "blocked_type")
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        if ext == "docx":
            return "docx"
        raise ApiError(400, "Only .docx Office files are supported.", "unsupported_type")
    if _image_mime(data):
        if not ocr_available():
            raise ApiError(400, "Image documents need OCR, which isn't installed on the server (tesseract).", "ocr_unavailable")
        return "image"
    if ext in ("pdf", "docx"):
        raise ApiError(400, f"This file isn't a valid {ext.upper()} file.", "bad_file")
    if ext in TEXT_EXT:
        if b"\x00" in data[:4096]:
            raise ApiError(400, "This file looks binary, not text.", "bad_file")
        return "text"
    raise ApiError(400, "Unsupported file type. Upload " + ", ".join(e.upper() for e in allowed_extensions()) + " files.",
                   "unsupported_type")


def _image_mime(data):
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[:4] in (b"II*\x00", b"MM\x00*"):
        return "image/tiff"
    return None


def _ocr_image(img):
    import pytesseract
    try:
        return pytesseract.image_to_string(img, timeout=60) or ""
    except RuntimeError:  # tesseract timeout
        return ""


def _check_docx(data):
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ApiError(400, "This DOCX file is damaged.", "bad_file")
    total = sum(i.file_size for i in zf.infolist())
    if total > 200 * 1024 * 1024 or len(zf.infolist()) > 5000:
        raise ApiError(400, "This DOCX file expands to an unsafe size.", "bad_file")
    if "word/document.xml" not in zf.namelist():
        raise ApiError(400, "This isn't a Word document.", "bad_file")


def extract_parts(kind, data):
    """Return ([(page_number, text)], used_ocr)."""
    try:
        if kind == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ApiError(400, "This PDF is password-protected.", "encrypted")
            if len(reader.pages) > settings.MAX_PDF_PAGES:
                raise ApiError(400, f"PDFs can have up to {settings.MAX_PDF_PAGES} pages.", "too_many_pages")
            parts = [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
            used_ocr = False
            empty = [i for i, (_, t) in enumerate(parts) if len(t.strip()) < 20]
            if empty and ocr_available() and importlib.util.find_spec("pypdfium2"):
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(data)
                for i in empty[:settings.OCR_MAX_PAGES]:
                    img = pdf[i].render(scale=2).to_pil()
                    txt = _ocr_image(img)
                    if txt.strip():
                        parts[i] = (i + 1, txt)
                        used_ocr = True
                pdf.close()
            return parts, used_ocr
        if kind == "docx":
            _check_docx(data)
            from docx import Document
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return [(1, "\n".join(parts))], False
        if kind == "image":
            from PIL import Image
            Image.MAX_IMAGE_PIXELS = 60_000_000  # decompression-bomb guard
            img = Image.open(io.BytesIO(data))
            frames = []
            for n in range(min(getattr(img, "n_frames", 1), settings.OCR_MAX_PAGES)):
                img.seek(n)
                frames.append((n + 1, _ocr_image(img.convert("RGB"))))
            return frames, True
        for enc in ("utf-8-sig", "cp1252"):
            try:
                return [(1, data.decode(enc))], False
            except UnicodeDecodeError:
                continue
        return [(1, data.decode("latin-1"))], False
    except ApiError:
        raise
    except ImportError:
        raise ApiError(500, "Document support is not fully installed on the server. Run: pip install -r requirements.txt",
                       "missing_dependency")
    except Exception:  # noqa: BLE001
        log.exception("Document parse failed")
        raise ApiError(400, "Could not read this file. It may be corrupted.", "bad_file")


def build_chunks(kind, parts):
    """Return ([(page, body)], page_count). PDFs/images keep real pages; text is grouped into ~3000-char sections."""
    if kind in ("pdf", "image"):
        return [(page, body) for page, text in parts for _, body in split_text(text)], len(parts)
    pieces = split_text(parts[0][1])
    return [(start // 3000 + 1, body) for start, body in pieces], (pieces[-1][0] // 3000 + 1 if pieces else 0)


# ------------------------------------------------------------------ embeddings
def pack(vec):
    return array.array("f", vec).tobytes()


def unpack(blob):
    a = array.array("f")
    a.frombytes(blob)
    return a


def cosine(a, b):
    dot = sum(x * y for x, y in zip(a, b, strict=False))  # vectors from different models may differ in length
    na = math.sqrt(sum(x * x for x in a)) or 1.0
    nb = math.sqrt(sum(y * y for y in b)) or 1.0
    return dot / (na * nb)


def embed_document(doc_id, uid):
    """Compute and store embeddings for a document's chunks. Returns True if embedded."""
    if not settings.EMBEDDING_MODEL:
        return False
    with db.engine.connect() as c:
        rows = c.execute(select(db.doc_chunks.c.id, db.doc_chunks.c.body).where(
            db.doc_chunks.c.document_id == doc_id, db.doc_chunks.c.user_id == uid).order_by(db.doc_chunks.c.idx)).all()
    if not rows:
        return False
    try:
        vectors = llm.embed([r.body for r in rows], {"uid": uid})
    except Exception as e:  # noqa: BLE001
        log.warning("embedding failed for %s: %s", doc_id, llm.friendly_error(e))
        return False
    with db.engine.begin() as c:
        for r, v in zip(rows, vectors, strict=True):
            c.execute(update(db.doc_chunks).where(db.doc_chunks.c.id == r.id).values(embedding=pack(v)))
        c.execute(update(db.documents).where(db.documents.c.id == doc_id).values(embedded=1))
    return True


# ------------------------------------------------------------------ ingestion
def ingest(uid, filename, data):
    """Validate, parse, chunk, deduplicate/version and store a document. Returns the new document row dict."""
    name = re.sub(r"[\x00-\x1f\\/]", "_", filename).strip()[:255] or "document"
    if not data:
        raise ApiError(400, "That file is empty.", "empty")
    if len(data) > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise ApiError(413, f"File is too large (max {settings.MAX_UPLOAD_MB} MB).", "too_large")
    kind = detect_kind(name, data)
    digest = hashlib.sha256(data).hexdigest()
    with db.engine.connect() as c:
        dup = c.execute(select(db.documents.c.filename).where(db.documents.c.user_id == uid,
                                                              db.documents.c.sha256 == digest)).first()
        if dup:
            raise ApiError(409, f"You already uploaded this file (as “{dup.filename}”).", "duplicate")
        owned = c.execute(select(func.count()).select_from(db.documents).where(db.documents.c.user_id == uid)).scalar()
    if owned >= settings.MAX_DOCS_PER_USER:
        raise ApiError(400, f"You can store up to {settings.MAX_DOCS_PER_USER} documents. Delete one first.", "doc_limit")
    parts, used_ocr = extract_parts(kind, data)
    pieces, page_count = build_chunks(kind, parts)
    if not pieces:
        hint = "" if ocr_available() else " Scanned or image-only files need OCR, which isn't installed on this server."
        raise ApiError(400, "No readable text found in this file." + hint, "no_text")
    store_kind = kind  # pdf | docx | text | image
    with db.engine.begin() as c:
        prev = c.execute(select(func.max(db.documents.c.version)).where(
            db.documents.c.user_id == uid, db.documents.c.filename == name)).scalar()
        doc = {"id": str(uuid.uuid4()), "user_id": uid, "filename": name, "kind": store_kind, "pages": page_count,
               "chunk_count": len(pieces), "size": len(data), "sha256": digest, "version": (prev or 0) + 1,
               "ocr": 1 if used_ocr else 0, "embedded": 0, "created_at": db.now_us()}
        c.execute(insert(db.documents).values(**doc))
        c.execute(insert(db.doc_chunks), [{"id": str(uuid.uuid4()), "document_id": doc["id"], "user_id": uid,
                                           "page": p, "idx": i, "body": b} for i, (p, b) in enumerate(pieces)])
    doc["embedded"] = 1 if embed_document(doc["id"], uid) else 0
    return doc


def reindex(uid, doc_id):
    """Recompute chunk embeddings for one document (e.g. after enabling or changing EMBEDDING_MODEL)."""
    with db.engine.connect() as c:
        if not c.execute(select(db.documents.c.id).where(db.documents.c.id == doc_id,
                                                         db.documents.c.user_id == uid)).first():
            raise ApiError(404, "Document not found.", "not_found")
    if not settings.EMBEDDING_MODEL:
        raise ApiError(400, "Semantic search is off (EMBEDDING_MODEL is not set). Keyword search needs no indexing.",
                       "embeddings_off")
    with db.engine.begin() as c:
        c.execute(update(db.doc_chunks).where(db.doc_chunks.c.document_id == doc_id).values(embedding=None))
        c.execute(update(db.documents).where(db.documents.c.id == doc_id).values(embedded=0))
    if not embed_document(doc_id, uid):
        raise ApiError(502, "Could not compute embeddings right now. Keyword search still works. Try again later.",
                       "embedding_failed")
    return True


def delete_document(uid, doc_id):
    with db.engine.begin() as c:
        if not c.execute(select(db.documents.c.id).where(db.documents.c.id == doc_id,
                                                         db.documents.c.user_id == uid)).first():
            raise ApiError(404, "Document not found.", "not_found")
        c.execute(delete(db.doc_chunks).where(db.doc_chunks.c.document_id == doc_id, db.doc_chunks.c.user_id == uid))
        c.execute(delete(db.documents).where(db.documents.c.id == doc_id, db.documents.c.user_id == uid))


# ------------------------------------------------------------------ retrieval
def retrieve(uid, doc_ids, query, k=None):
    """Top-k chunks from the user's selected documents. Ownership is enforced in the query itself."""
    k = k or settings.RAG_TOP_K
    if not doc_ids:
        return []
    with db.engine.connect() as c:
        rows = c.execute(
            select(db.doc_chunks.c.document_id, db.doc_chunks.c.page, db.doc_chunks.c.body, db.doc_chunks.c.embedding,
                   db.documents.c.filename, db.documents.c.kind)
            .select_from(db.doc_chunks.join(db.documents, db.doc_chunks.c.document_id == db.documents.c.id))
            .where(db.doc_chunks.c.user_id == uid, db.documents.c.user_id == uid, db.doc_chunks.c.document_id.in_(doc_ids))
            .order_by(db.documents.c.created_at, db.doc_chunks.c.idx)).all()
    if not rows:
        return []
    kw = bm25_scores(query, [r.body for r in rows])
    kw_rank = [i for i in sorted(range(len(rows)), key=lambda i: -kw[i]) if kw[i] > 0]
    vec_rank = []
    if settings.EMBEDDING_MODEL and any(r.embedding for r in rows):
        try:
            qv = llm.embed([query], {"uid": uid})[0]
            sims = [(cosine(qv, unpack(r.embedding)) if r.embedding else -1.0) for r in rows]
            vec_rank = [i for i in sorted(range(len(rows)), key=lambda i: -sims[i]) if sims[i] > 0.2]
        except Exception as e:  # noqa: BLE001  fall back to keyword search
            log.warning("query embedding failed, using keyword search only: %s", llm.friendly_error(e))
    if vec_rank:
        fused = Counter()
        for rank_list in (kw_rank, vec_rank):
            for pos, i in enumerate(rank_list[:50]):
                fused[i] += 1.0 / (60 + pos)
        order = [i for i, _ in fused.most_common(k)]
    else:
        order = kw_rank[:k]
    if not order:  # broad requests like "summarize this" share no keywords: sample across the documents
        order = list(range(0, len(rows), max(1, len(rows) // k)))[:k]
    return [rows[i] for i in order]
